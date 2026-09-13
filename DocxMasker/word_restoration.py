# word_restoration.py
import json
from docx import Document

def Word_restoration(
    modified_file: str, json_file_path: str, output_restored_file: str
) -> bool:
    """
    Restore encrypted tokens in DOCX file to original text using JSON mapping dict.
    """
    try:
        with open(json_file_path, "r", encoding="utf-8") as f:
            mapping = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError) as e:
        print(f"⚠️ JSON mapping file error: {e}")
        return False

    try:
        content = Document(modified_file)
    except FileNotFoundError:
        print(f"⚠️ Modified DOCX file not found: {modified_file}")
        return False
    except Exception as e:
        print(f"⚠️ Failed to open DOCX file: {e}")
        return False

    try:
        #ترتيب الكلمات المشفرة من الأطول إلى الأقصر
        sorted_tokens = sorted(mapping.keys(), key=len, reverse=True)

        def replace_in_paragraphs(paragraphs):
            for p in paragraphs:
                if not p.text:
                    continue

                original_text = p.text
                is_modified = False

                # استخدام القائمة المرتبة بدلاً من القاموس العشوائي
                for new_word in sorted_tokens:
                    if new_word in original_text:
                        original_text = original_text.replace(new_word, mapping[new_word])
                        is_modified = True

                if is_modified:
                    p.text = original_text

        def replace_in_table(table):
            for row in table.rows:
                for cell in row.cells:
                    replace_in_paragraphs(cell.paragraphs)
                    for nested_table in cell.tables:
                        replace_in_table(nested_table)

        replace_in_paragraphs(content.paragraphs)

        for table in content.tables:
            replace_in_table(table)

        for section in content.sections:
            for attr in ("header", "footer"):
                part = getattr(section, attr)
                replace_in_paragraphs(part.paragraphs)
                for table in part.tables:
                    replace_in_table(table)

        content.save(output_restored_file)
        return True
    except Exception as e:
        print(f"⚠️ Restoration process error: {e}")
        return False
